"""Generate sample_kb/summit-customer-handbook-DEMO.docx: a longer, realistic Word document with
real heading levels, long paragraphs and a table, so chunking strategies actually differ.
DEMO CONTENT - fictional policies for testing. Replace with the contractor's real handbook."""
import docx

H = "heading"
SECTIONS = [
    (H, 1, "Summit Roofing & Exteriors Customer Handbook (DEMO)"),
    ("p", "This handbook explains what homeowners can expect before, during and after a roofing project with "
          "Summit Roofing & Exteriors. It covers preparation, scheduling, warranties, payment, insurance claims "
          "and common questions. It is provided for demonstration purposes and describes fictional policies."),

    (H, 1, "Before Your Project"),
    (H, 2, "Permits"),
    ("p", "Most full roof replacements require a building permit from the local municipality. Summit applies for "
          "the permit on the homeowner's behalf once the contract is signed, and the permit fee is included in "
          "the project price. Repairs smaller than one roofing square usually do not need a permit. Permit "
          "approval typically takes five to ten business days, and work cannot begin until the permit is issued."),
    (H, 2, "Preparing Your Home"),
    ("p", "Please move vehicles out of the driveway and garage the evening before work begins, because shingle "
          "deliveries and the waste bin are usually placed on the driveway. Remove fragile items from walls and "
          "shelves, since hammering on the roof causes vibration inside the house. Cover belongings stored in the "
          "attic with old sheets; some dust and small debris can fall through the roof boards when old shingles "
          "are removed. Patio furniture, planters and barbecues should be moved at least ten feet from the house."),
    ("p", "Pets should be kept indoors or away from the property during working hours. The noise can be "
          "stressful for animals, and open gates and ladders create a risk of escape or injury."),
    (H, 2, "Choosing Colours and Materials"),
    ("p", "Homeowners choose shingle colours from physical samples brought to the site visit. Because available "
          "colours depend on distributor inventory, Summit confirms the colour with the supplier before ordering "
          "and will call the homeowner if a colour is unavailable. Changing the colour after materials are "
          "ordered may delay the project by up to two weeks."),

    (H, 1, "Scheduling and Weather"),
    (H, 2, "Lead Times"),
    ("p", "During spring and early summer, the waiting time between signing a contract and starting work is "
          "usually three to five weeks. In late summer and autumn it is usually two to three weeks. Emergency "
          "repairs and leak assessments are scheduled separately and are not affected by these lead times."),
    (H, 2, "Weather Delays"),
    ("p", "Roofing is not carried out in rain, high winds above 40 kilometres per hour, or when temperatures are "
          "below minus five degrees Celsius, because shingles do not seal properly in the cold. If weather "
          "forces a delay, the project coordinator contacts the homeowner by phone or text the same morning and "
          "offers the next available day. A roof is never left open overnight: if weather changes during the "
          "day, the crew installs underlayment and tarps before leaving."),
    (H, 2, "How Long a Replacement Takes"),
    ("p", "A typical single-family home is completed in one to two days. Larger homes, steep roofs, multiple "
          "layers of old shingles or replacement of rotten decking can extend the work to three or four days. "
          "The crew arrives at about 7:30 AM and normally finishes by 6:00 PM."),

    (H, 1, "During the Work"),
    (H, 2, "Decking Replacement"),
    ("p", "When the old shingles are removed, the crew inspects the wooden roof decking underneath. Soft or "
          "rotten boards must be replaced so that new shingles can be fastened securely. Decking replacement "
          "is not included in the replacement price per square, because its condition cannot be seen until the "
          "old roof is removed. It is charged per sheet of plywood replaced, and the homeowner is shown the "
          "damaged boards before the work is done."),
    (H, 2, "Ventilation and Ice Protection"),
    ("p", "Every replacement includes ice and water shield membrane along the eaves and in valleys, and a "
          "ventilation review. Poor attic ventilation shortens shingle life and can void manufacturer "
          "warranties, so Summit may recommend additional roof vents or a ridge vent. Any recommended "
          "ventilation upgrade is quoted before work begins."),
    (H, 2, "Clean-up"),
    ("p", "At the end of every working day the crew cleans the site. On the final day the lawn, driveway and "
          "flower beds are swept with a magnetic roller to collect nails. The waste bin is removed within two "
          "business days of completion."),

    (H, 1, "Warranties"),
    (H, 2, "Workmanship Warranty"),
    ("p", "Summit provides a ten-year workmanship warranty on full roof replacements and a two-year workmanship "
          "warranty on repairs. The workmanship warranty covers defects in installation, such as leaks caused "
          "by improperly installed flashing or shingles. It does not cover damage from storms, falling trees, "
          "ice dams caused by poor insulation, or work carried out by other contractors."),
    (H, 2, "Manufacturer Warranty"),
    ("p", "Shingle manufacturers provide their own warranties on materials. GAF Timberline HDZ and CertainTeed "
          "Landmark shingles carry manufacturer warranties against manufacturing defects; the length and terms "
          "are set by the manufacturer and are explained in the written proposal. Manufacturer warranties may "
          "require the home to have adequate attic ventilation."),
    (H, 2, "Making a Warranty Claim"),
    ("p", "To make a warranty claim, contact the office by phone or email with the project address and a "
          "description of the problem. A roofing specialist will review the claim and arrange an inspection, "
          "normally within five business days. If the problem is an active leak, it is handled as an emergency "
          "leak enquiry. Warranty inspections for covered defects are free of charge."),
    (H, 2, "Transferring the Warranty"),
    ("p", "The workmanship warranty on a full replacement can be transferred once to a new owner if the house "
          "is sold within the warranty period. The new owner must register the transfer with Summit within "
          "sixty days of the sale."),

    (H, 1, "Payment"),
    ("p", "Payment terms for roof replacements are shown below. Repairs and inspections are paid in full on "
          "completion."),
    ("table", [["Stage", "Amount", "When"],
               ["Deposit", "30% of the contract price", "When the contract is signed"],
               ["Balance", "Remaining 70%", "On completion, after the final walk-through"]]),
    ("p", "Summit accepts e-transfer, cheque, debit and credit card. Credit card payments over 5,000 dollars "
          "carry a 2.5% processing fee. Financing is not offered directly, but some customers use a home "
          "improvement line of credit from their bank."),

    (H, 1, "Insurance Claims"),
    ("p", "If a homeowner believes storm or hail damage may be covered by insurance, Summit can carry out an "
          "inspection and provide a written report with photographs that the homeowner can submit to the "
          "insurer. Summit does not negotiate with insurance companies on the homeowner's behalf and cannot "
          "say whether a claim will be approved. The homeowner should contact their insurer before any repair "
          "work begins, except for temporary measures to stop an active leak."),

    (H, 1, "Frequently Asked Questions"),
    (H, 2, "Do I need to be home during the work?"),
    ("p", "No. The homeowner needs to be available by phone, and the crew needs access to an outdoor power "
          "outlet. Someone should be home for the final walk-through if possible."),
    (H, 2, "Can you roof over my existing shingles?"),
    ("p", "No. Summit always removes the old roofing down to the decking. Installing new shingles over old ones "
          "hides decking damage, adds weight, and is not permitted by most manufacturer warranties."),
    (H, 2, "Will my gutters be affected?"),
    ("p", "Gutters stay in place during a replacement unless they are being replaced at the same time. If "
          "existing gutters are damaged during the work, Summit repairs them at no cost."),
    (H, 2, "Do you do skylights or chimneys?"),
    ("p", "Summit replaces flashing around existing skylights and chimneys as part of a roof replacement. "
          "Installing new skylights and chimney masonry repairs are referred to specialist partners."),
]


def build(path: str) -> None:
    d = docx.Document()
    for item in SECTIONS:
        if item[0] == H:
            d.add_heading(item[2], level=item[1])
        elif item[0] == "p":
            d.add_paragraph(item[1])
        elif item[0] == "table":
            rows = item[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
    d.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    build("sample_kb/summit-customer-handbook-DEMO.docx")
