import dataclasses
import io
from pathlib import Path

import docx
from openpyxl import Workbook

from app import chunking, knowledge, vectorstore
from app.chunking import ChunkConfig
from app.config import settings

SAMPLE = Path(__file__).resolve().parent.parent / "sample_kb"


def test_plain_text_headings_detected():
    text = knowledge.extract_text("info.txt", (SAMPLE / "summit-company-info.txt").read_bytes())
    headings = [ln for ln in text.splitlines() if ln.startswith("# ")]
    assert "# Emergency Leak Requests" in headings
    assert "# GAF Timberline HDZ" in headings
    assert not any("Sunday" in h or "Monday" in h for h in headings)  # "Sunday: Closed" is not a heading


def test_chunks_carry_contextual_header_and_never_split_lines():
    text = "# Price List\n" + "\n".join(f"- Item {i}: some long description of the item; Price: {i * 10}" for i in range(40))
    chunks = knowledge.chunk_text(text, "prices.xlsx", ChunkConfig(max_chars=400))
    assert len(chunks) > 1
    for c in chunks:
        assert c["text"].startswith("[prices.xlsx > Price List]")
        for line in c["text"].splitlines()[1:]:
            assert line.startswith("- Item") and "Price:" in line  # rows stay whole


def test_docx_keeps_heading_hierarchy_and_tables():
    d = docx.Document()
    d.add_heading("Services", level=1)
    d.add_heading("Gutters", level=2)
    d.add_paragraph("We replace gutters.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Downspout", "120"
    buf = io.BytesIO(); d.save(buf)
    text = knowledge.extract_text("services.docx", buf.getvalue())
    assert "# Services" in text and "## Gutters" in text and "- Item: Downspout; Price: 120" in text
    chunks = knowledge.chunk_text(text, "services.docx")
    assert chunks[-1]["text"].startswith("[services.docx > Services > Gutters]")


def test_xlsx_rows_are_self_describing():
    wb = Workbook(); ws = wb.active; ws.title = "Prices"
    ws.append(["Item", "Unit", "Price"]); ws.append(["Gutter", "per ft", 14])
    buf = io.BytesIO(); wb.save(buf)
    text = knowledge.extract_text("p.xlsx", buf.getvalue())
    assert "# Prices" in text and "- Item: Gutter; Unit: per ft; Price: 14" in text


def test_small_kb_uses_full_context(contractor):
    knowledge.add_document(contractor["id"], "info.txt", "# About\nWe fix roofs.")
    status = knowledge.kb_status(contractor["id"])
    assert status["mode"] == "full_context" and status["backend"] == "prompt"


def test_crossing_threshold_triggers_pinecone_sync(contractor, monkeypatch):
    monkeypatch.setattr(knowledge, "settings", dataclasses.replace(settings, kb_full_context_max_tokens=50))
    upserted, deleted = [], []
    monkeypatch.setattr(vectorstore, "enabled", lambda: True)
    monkeypatch.setattr(vectorstore, "upsert", lambda cid, rows: upserted.extend(r["uid"] for r in rows))
    monkeypatch.setattr(vectorstore, "delete", lambda cid, uids: deleted.extend(uids))

    small = knowledge.add_document(contractor["id"], "a.txt", "# About\nWe fix roofs.")
    assert small["mode"] == "full_context" and upserted == []          # under threshold: nothing synced

    big = knowledge.add_document(contractor["id"], "b.txt", "# Catalog\n" + "Shingle spec line. " * 40)
    assert big["mode"] == "retrieval" and big["backend"] == "pinecone"
    assert big["unsynced_chunks"] == 0
    assert any(u.startswith(f"c{contractor['id']}-d") for u in upserted)
    assert len(upserted) == len(set(upserted))                          # earlier doc synced too, once

    doc_id = big["document_id"]
    knowledge.delete_document(contractor["id"], doc_id)
    assert deleted and all(f"-d{doc_id}-" in u for u in deleted)


def test_retrieval_falls_back_to_bm25_without_pinecone(contractor, monkeypatch):
    monkeypatch.setattr(knowledge, "settings", dataclasses.replace(settings, kb_full_context_max_tokens=10))
    text = knowledge.extract_text("info.txt", (SAMPLE / "summit-company-info.txt").read_bytes())
    knowledge.add_document(contractor["id"], "info.txt", text)
    status = knowledge.kb_status(contractor["id"])
    assert status["mode"] == "retrieval" and status["backend"] == "local-bm25"
    found = knowledge.search(contractor["id"], "GAF Timberline HDZ")
    assert found["backend"] == "local-bm25"
    assert "Timberline" in found["results"][0]["text"]


def test_tenants_are_isolated(contractor):
    from app import accounts
    other = accounts.create_contractor("o@example.com", "password1", "Other Co", "Sam")
    knowledge.add_document(other, "secret.txt", "# Prices\nSecret discount code ZEBRA")
    assert knowledge.bm25_search(contractor["id"], "ZEBRA discount") == []
