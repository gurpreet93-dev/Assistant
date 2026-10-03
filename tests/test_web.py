import io

import docx
from fastapi.testclient import TestClient

from app import agent
from app.main import app


def _login(contractor):
    c = TestClient(app)
    c.post("/login", data={"email": "t@example.com", "password": "password1"})
    return c


def test_upload_word_document_and_view_chunks(contractor):
    c = _login(contractor)
    d = docx.Document(); d.add_heading("Warranty", 1); d.add_paragraph("10-year workmanship warranty.")
    buf = io.BytesIO(); d.save(buf)
    r = c.post("/knowledge/upload", files={"files": ("warranty.docx", buf.getvalue(),
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200 and "warranty.docx" in r.text and "Full-context mode" in r.text


def test_unsupported_file_is_reported_not_crashing(contractor):
    c = _login(contractor)
    r = c.post("/knowledge/upload", files={"files": ("photo.png", b"\x89PNG", "image/png")})
    assert r.status_code == 200 and "Unsupported file type" in r.text


def test_pages_require_login():
    r = TestClient(app).get("/dashboard", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"


def test_twilio_flow_greets_listens_and_hangs_up(contractor, monkeypatch):
    c = TestClient(app)
    r = c.post("/voice/incoming", data={"CallSid": "CA1", "From": "+15555550123", "To": "+14165550100"})
    assert "<Gather" in r.text and 'input="speech"' in r.text and "Summit Roofing &amp; Exteriors" in r.text

    r = c.post("/voice/turn", data={"CallSid": "CA1", "SpeechResult": ""})
    assert "didn't catch that" in r.text

    monkeypatch.setattr(agent, "handle_turn", lambda sid, text: {"say": "Goodbye!", "ended": True, "actions": []})
    r = c.post("/voice/turn", data={"CallSid": "CA1", "SpeechResult": "that's all"})
    assert "<Hangup/>" in r.text and "Goodbye!" in r.text

    assert c.post("/voice/status", data={"CallSid": "CA1", "CallStatus": "completed"}).status_code == 204
    assert agent.get_call("CA1")["status"] == "completed"
