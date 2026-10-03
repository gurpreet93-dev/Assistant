"""Knowledge base: ingest contractor documents and decide how the agent will read them.

Routing (decision log D6):
  * total KB <= KB_FULL_CONTEXT_MAX_TOKENS  -> "full_context": every document goes into the
    system prompt (prompt-cached). No retrieval, so nothing can be missed.
  * total KB  > threshold                   -> "retrieval": chunks are synced to Pinecone
    (one namespace per contractor) and the agent gets a search tool.

Chunks are always computed and stored in SQLite, so switching modes never needs a re-upload.
If Pinecone isn't configured or fails, retrieval falls back to local BM25 keyword search.
"""
import io
import logging
import math
import re
from collections import Counter

from app import vectorstore
from app.config import settings
from app.db import get_db, now_iso

log = logging.getLogger(__name__)

CHUNK_CHARS = 800
SUPPORTED = (".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv")
STOPWORDS = set(
    "a an and are as at be by do does for from has have how i in is it its me my of on or "
    "our so that the their them they this to us was we what when where which who will with "
    "you your can if any".split()
)


# ---------------------------------------------------------------- extraction
# Every format is normalised to plain text with markdown '#' headings, so one chunker serves all.

def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".xlsx"):
        return _extract_xlsx(data)
    if name.endswith((".txt", ".md", ".csv")):
        return _mark_plain_headings(data.decode("utf-8", errors="replace"))
    raise ValueError(f"Unsupported file type. Upload one of: {', '.join(SUPPORTED)}")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return _mark_plain_headings("\n".join((page.extract_text() or "") for page in reader.pages))


def _extract_docx(data: bytes) -> str:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    lines = []
    for block in document.iter_inner_content():  # paragraphs and tables, in document order
        if isinstance(block, Table):
            rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
            if rows:
                header = rows[0]
                for row in rows[1:]:
                    lines.append("- " + "; ".join(f"{h}: {v}" for h, v in zip(header, row) if v))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name or "").lower() if block.style is not None else ""
        if style.startswith("heading") or style == "title":
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            lines.append(f"{'#' * level} {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n".join(lines)


def _extract_xlsx(data: bytes) -> str:
    """Each sheet becomes a section; each row becomes one self-describing line
    ('Item: Gutter replacement; Unit: linear ft; Price: 14'), so a row is never split
    from its column names."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        rows = [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        lines.append(f"# {ws.title}")
        header = rows[0]
        for row in rows[1:]:
            lines.append("- " + "; ".join(f"{h or f'Column {i + 1}'}: {v}" for i, (h, v) in enumerate(zip(header, row)) if v))
    return "\n".join(lines)


def _looks_like_heading(line: str, next_line: str | None) -> bool:
    """Plain text / PDF has no styles, so guess: short, title-like, no digits, no
    colon, no closing punctuation, and followed by more text."""
    if next_line is None or not line or line[0] in "-*•" or not line[0].isupper():
        return False
    return (len(line) <= 60 and len(line.split()) <= 8 and not re.search(r"[\d:]", line)
            and line[-1] not in ".,;!?)")


def _mark_plain_headings(text: str) -> str:
    raw = [ln.strip() for ln in text.splitlines()]
    if any(ln.startswith("#") for ln in raw):  # already markdown
        return "\n".join(raw)
    out = []
    for i, line in enumerate(raw):
        nxt = next((ln for ln in raw[i + 1:] if ln), None)
        out.append(f"# {line}" if _looks_like_heading(line, nxt) else line)
    return "\n".join(out)


# ---------------------------------------------------------------- chunking (decision log D4)

def chunk_text(text: str, source: str, max_chars: int = CHUNK_CHARS) -> list[dict]:
    """Split at headings, pack whole lines up to max_chars (a line - a price row, a bullet -
    is never split), and prefix each chunk with '[source > heading path]' so it still makes
    sense on its own after retrieval."""
    sections: list[tuple[str, list[str]]] = []
    stack: list[tuple[int, str]] = []
    current: list[str] = []

    def flush():
        if any(ln.strip() for ln in current):
            sections.append((" > ".join(h for _, h in stack), [ln for ln in current if ln.strip()]))

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            flush()
            current = []
            level = len(m.group(1))
            stack = [(lvl, h) for lvl, h in stack if lvl < level] + [(level, m.group(2).strip())]
        else:
            current.append(line)
    flush()

    chunks = []
    for path, lines in sections:
        header = f"[{source}{' > ' + path if path else ''}]"
        body = ""
        for line in lines:
            if body and len(header) + len(body) + len(line) + 2 > max_chars:
                chunks.append({"section": path, "text": f"{header}\n{body}"})
                body = line
            else:
                body = f"{body}\n{line}" if body else line
        if body:
            chunks.append({"section": path, "text": f"{header}\n{body}"})
    return chunks


def estimate_tokens(text: str) -> int:
    # ~4 characters per token for English. Good enough for a routing threshold;
    # use the count_tokens API if the threshold ever needs to be exact.
    return math.ceil(len(text) / 4)


# ---------------------------------------------------------------- storage + routing

def add_document(contractor_id: int, filename: str, content: str) -> dict:
    chunks = chunk_text(content, filename)
    with get_db() as db:
        doc_id = db.execute(
            "INSERT INTO documents (contractor_id, filename, content, token_estimate, created_at) VALUES (?,?,?,?,?)",
            (contractor_id, filename, content, estimate_tokens(content), now_iso()),
        ).lastrowid
        for i, chunk in enumerate(chunks):
            db.execute(
                "INSERT INTO chunks (document_id, contractor_id, idx, uid, section, text) VALUES (?,?,?,?,?,?)",
                (doc_id, contractor_id, i, f"c{contractor_id}-d{doc_id}-{i}", chunk["section"], chunk["text"]),
            )
    return {"document_id": doc_id, "chunks": len(chunks), **sync_mode(contractor_id)}


def delete_document(contractor_id: int, doc_id: int) -> dict:
    with get_db() as db:
        synced = [r["uid"] for r in db.execute(
            "SELECT uid FROM chunks WHERE document_id=? AND contractor_id=? AND synced=1", (doc_id, contractor_id))]
        db.execute("DELETE FROM documents WHERE id=? AND contractor_id=?", (doc_id, contractor_id))
    if synced and vectorstore.enabled():
        try:
            vectorstore.delete(contractor_id, synced)
        except Exception:  # stale vectors are harmless: search re-checks chunk ids against SQLite
            log.exception("Pinecone delete failed")
    return sync_mode(contractor_id)


def kb_status(contractor_id: int) -> dict:
    with get_db() as db:
        row = db.execute(
            "SELECT count(*) AS docs, coalesce(sum(token_estimate),0) AS tokens FROM documents WHERE contractor_id=?",
            (contractor_id,),
        ).fetchone()
        unsynced = db.execute("SELECT count(*) FROM chunks WHERE contractor_id=? AND synced=0",
                              (contractor_id,)).fetchone()[0]
    mode = "full_context" if row["tokens"] <= settings.kb_full_context_max_tokens else "retrieval"
    return {"mode": mode, "documents": row["docs"], "tokens": row["tokens"],
            "threshold": settings.kb_full_context_max_tokens, "unsynced_chunks": unsynced,
            "backend": ("pinecone" if vectorstore.enabled() else "local-bm25") if mode == "retrieval" else "prompt"}


def sync_mode(contractor_id: int) -> dict:
    """The automatic trigger: once the KB crosses the threshold, push every not-yet-synced
    chunk to Pinecone. Runs after every upload/delete."""
    status = kb_status(contractor_id)
    if status["mode"] != "retrieval" or not status["unsynced_chunks"] or not vectorstore.enabled():
        return status
    with get_db() as db:
        rows = [dict(r) for r in db.execute(
            "SELECT c.uid, c.text, c.section, c.document_id, d.filename FROM chunks c "
            "JOIN documents d ON d.id=c.document_id WHERE c.contractor_id=? AND c.synced=0", (contractor_id,))]
    error = None
    try:
        vectorstore.upsert(contractor_id, rows)
        with get_db() as db:
            db.executemany("UPDATE chunks SET synced=1 WHERE uid=?", [(r["uid"],) for r in rows])
    except Exception as exc:
        log.exception("Pinecone sync failed")
        error = str(exc)
    status = kb_status(contractor_id)
    if error:
        status["sync_error"] = error
    return status


def list_documents(contractor_id: int) -> list[dict]:
    with get_db() as db:
        rows = db.execute(
            "SELECT d.id, d.filename, d.created_at, d.token_estimate, "
            "(SELECT count(*) FROM chunks c WHERE c.document_id=d.id) AS chunk_count, "
            "(SELECT count(*) FROM chunks c WHERE c.document_id=d.id AND c.synced=1) AS synced_count "
            "FROM documents d WHERE d.contractor_id=? ORDER BY d.created_at DESC, d.id DESC",
            (contractor_id,),
        ).fetchall()
    return [dict(r) for r in rows]


def full_text(contractor_id: int) -> list[dict]:
    """All documents, for the full-context prompt. Ordered by id so the prompt is byte-stable
    across turns (any change would invalidate the prompt cache)."""
    with get_db() as db:
        rows = db.execute("SELECT filename, content FROM documents WHERE contractor_id=? ORDER BY id",
                          (contractor_id,)).fetchall()
    return [dict(r) for r in rows]


# ---------------------------------------------------------------- search (retrieval mode)

def search(contractor_id: int, query: str, k: int = 4) -> dict:
    if vectorstore.enabled():
        try:
            hits = vectorstore.search(contractor_id, query, k)
            with get_db() as db:  # drop hits for chunks deleted since they were indexed
                live = {r["uid"] for r in db.execute("SELECT uid FROM chunks WHERE contractor_id=?", (contractor_id,))}
            return {"backend": "pinecone", "results": [h for h in hits if h["id"] in live]}
        except Exception as exc:
            log.exception("Pinecone search failed, falling back to BM25")
            return {"backend": "local-bm25", "fallback_reason": str(exc), "results": bm25_search(contractor_id, query, k)}
    return {"backend": "local-bm25", "results": bm25_search(contractor_id, query, k)}


def _tokenize(text: str) -> list[str]:
    return [t for t in re.findall(r"[a-z0-9]+", text.lower()) if t not in STOPWORDS]


def bm25_search(contractor_id: int, query: str, k: int = 4) -> list[dict]:
    """Keyword ranking (BM25, k1=1.5, b=0.75). Fallback when Pinecone is unavailable,
    and strong on exact product names where embeddings are weak."""
    with get_db() as db:
        rows = db.execute("SELECT c.uid, c.text, d.filename FROM chunks c JOIN documents d ON d.id=c.document_id "
                          "WHERE c.contractor_id=?", (contractor_id,)).fetchall()
    if not rows:
        return []
    docs = [_tokenize(r["text"]) for r in rows]
    n, avgdl = len(docs), (sum(len(d) for d in docs) / len(docs)) or 1
    df = Counter(term for d in docs for term in set(d))
    scored = []
    for row, doc in zip(rows, docs):
        tf, score = Counter(doc), 0.0
        for term in _tokenize(query):
            if term in tf:
                idf = math.log(1 + (n - df[term] + 0.5) / (df[term] + 0.5))
                score += idf * tf[term] * 2.5 / (tf[term] + 1.5 * (0.25 + 0.75 * len(doc) / avgdl))
        if score > 0:
            scored.append({"id": row["uid"], "source": row["filename"], "text": row["text"], "score": round(score, 2)})
    return sorted(scored, key=lambda r: r["score"], reverse=True)[:k]
